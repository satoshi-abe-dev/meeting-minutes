"""Convert minutes Markdown to .docx (Word).

Does not use a general-purpose Markdown parser. Converts only the Markdown
subset that minutes actually use (the built-in template, custom templates,
"Auto" mode) with a small line-based state machine. An unexpected line falls
back to a plain paragraph, and the conversion itself never raises (.docx is
a side product of minutes.md; the policy here is to never let it fail the
run).

Supported syntax:
    - ATX headings `#` through `######`   -> Word's Heading 1..6
    - Bullet lists `- ` / `* ` / `+ ` (nested by leading spaces) -> List Bullet / 2 / 3
    - Numbered lists `1. ` / `1) `        -> a plain paragraph that keeps the original number
      (Doesn't use Word's auto-numbered List Number: since it shares a
       numbering instance, a second numbered list separated by a heading
       would continue the previous numbering, and its start number would be
       lost. Minutes barely use numbered lists, so this simplification is
       enough)
    - GFM pipe tables (`| ... |` rows + a `| --- |` separator row) -> a table (header row cells bolded)
    - Fenced code blocks ``` ... ```      -> a monospace (Consolas) paragraph
    - Blank line -> paragraph break / any other line -> a plain paragraph
    - Inline: only `**bold**` and `` `code` `` (nesting not supported)
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:  # A heavy dependency; imported inside the function at runtime.
    from docx.document import Document

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\s*(\d+[.)])\s+(.*)$")
_TABLE_SEP_CELL_RE = re.compile(r"^\s*:?-+:?\s*$")
# `**bold**` (1+ chars) or `` `code` `` (1+ chars, no backtick inside).
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

_MONO_FONT = "Consolas"
_MAX_LIST_DEPTH = 3  # up to List Bullet / List Bullet 2 / List Bullet 3


# --- Inline ------------------------------------------------------------

def _iter_inline(text: str) -> Iterator[tuple[str, bool, bool]]:
    """Split text into a sequence of (fragment, is_bold, is_monospace_code)."""
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], False, False
        tok = m.group(0)
        if tok.startswith("**"):
            yield tok[2:-2], True, False
        else:  # `code`
            yield tok[1:-1], False, True
        pos = m.end()
    if pos < len(text):
        yield text[pos:], False, False


def _add_runs(paragraph, text: str, *, force_bold: bool = False) -> None:
    """Write text into paragraph, applying inline markup."""
    for seg, is_bold, is_code in _iter_inline(text):
        run = paragraph.add_run(seg)
        if is_bold or force_bold:
            run.bold = True
        if is_code:
            run.font.name = _MONO_FONT


# --- Tables ------------------------------------------------------------------

def _looks_like_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def _split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    cells = re.split(r"(?<!\\)\|", s)
    return [c.replace("\\|", "|").strip() for c in cells]


def _is_table_separator(line: str) -> bool:
    if not _looks_like_row(line):
        return False
    cells = _split_row(line)
    return bool(cells) and all(_TABLE_SEP_CELL_RE.match(c) for c in cells)


def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    ncols = max(1, len(header))
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    for k in range(ncols):
        cell = table.rows[0].cells[k]
        _add_runs(cell.paragraphs[0], header[k] if k < len(header) else "",
                  force_bold=True)
    for row in rows:
        cells = table.add_row().cells
        for k in range(ncols):
            _add_runs(cells[k].paragraphs[0], row[k] if k < len(row) else "")


# --- Code blocks -----------------------------------------------------

def _add_code_block(doc: Document, code_lines: list[str]) -> None:
    p = doc.add_paragraph(style="No Spacing")
    run = p.add_run("\n".join(code_lines))
    run.font.name = _MONO_FONT


# --- The conversion itself ----------------------------------------------------------

def markdown_to_docx(markdown: str) -> Document:
    """Convert minutes Markdown to a python-docx Document and return it.

    An unknown line falls back to a plain paragraph. Even if processing a
    line unexpectedly raises, that line is turned into a plain paragraph
    and skipped; the conversion as a whole never stops.
    """
    import docx

    doc = docx.Document()
    lines = markdown.splitlines()
    n = len(lines)
    i = 0
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Fenced code block
        if stripped.startswith("```"):
            j = i + 1
            buf: list[str] = []
            while j < n and not lines[j].strip().startswith("```"):
                buf.append(lines[j])
                j += 1
            _add_code_block(doc, buf)
            i = j + 1 if j < n else n
            continue

        if not stripped:
            i += 1
            continue

        try:
            # GFM pipe table (treated as a table only when a separator row follows)
            if (
                _looks_like_row(line)
                and i + 1 < n
                and _is_table_separator(lines[i + 1])
            ):
                header = _split_row(line)
                body: list[list[str]] = []
                j = i + 2
                while j < n and _looks_like_row(lines[j]):
                    body.append(_split_row(lines[j]))
                    j += 1
                _add_table(doc, header, body)
                i = j
                continue

            m = _HEADING_RE.match(line)
            if m:
                p = doc.add_paragraph(style=f"Heading {len(m.group(1))}")
                _add_runs(p, m.group(2).strip())
                i += 1
                continue

            m = _BULLET_RE.match(line)
            if m:
                depth = min(len(m.group(1)) // 2, _MAX_LIST_DEPTH - 1)
                style = "List Bullet" if depth == 0 else f"List Bullet {depth + 1}"
                _add_runs(doc.add_paragraph(style=style), m.group(2).strip())
                i += 1
                continue

            m = _NUMBER_RE.match(line)
            if m:
                # A plain paragraph that keeps the original Markdown number
                # (`1.` / `3)` etc.) literally. Doesn't use Word's
                # auto-numbering (List Number) — see the docstring.
                p = doc.add_paragraph()
                p.add_run(m.group(1) + " ")
                _add_runs(p, m.group(2).strip())
                i += 1
                continue

            _add_runs(doc.add_paragraph(), stripped)
        except Exception:
            # Don't let .docx generation stop even on an unexpected line.
            # Fall back to a plain paragraph and move on.
            doc.add_paragraph(stripped)
        i += 1

    return doc


def save_minutes_docx(markdown: str, out_dir: str | Path) -> Path:
    """Write minutes Markdown out as out_dir/minutes.docx and return its path."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "minutes.docx"
    markdown_to_docx(markdown).save(str(path))
    return path
